import json
import os
import shutil
import subprocess
from pathlib import Path

import pytest

from import_results import (
    Paths,
    ResultsError,
    Route,
    find_results,
    import_attempts,
    load_results,
    main,
    pending,
)

REAL_REGISTRAR = Path("~/Documents/TheMind/ATC/teoria/99-SCRIPTS/registrar.py").expanduser()
REAL_EXCEL = Path(
    "~/Documents/controlador/Primer Examen/Temario/Teoria_Maestro_v1.xlsx"
).expanduser()


def item(q, key, ans, *, doubt=False, at=None):
    return {
        "q": q,
        "key": key,
        "ans": ans,
        "doubt": doubt,
        "stuck": False,
        "secs": 50,
        "first_answered_at_secs": at,
        "cause": None,
    }


def attempt(uuid="u1", test_id="SYN-T9-001", number=1, started="2026-10-14T20:30:00+02:00"):
    return {
        "uuid": uuid,
        "test_id": test_id,
        "kind": "fixed",
        "mode": "exam",
        "attempt": number,
        "started": started,
        "finished": "2026-10-14T20:32:00+02:00",
        "limit_min": 2,
        "total_secs": 100,
        "status": "finished",
        "items": [
            item(f"{test_id}#1", "c", "c", at=50),
            item(f"{test_id}#2", "a", None, doubt=True),
        ],
    }


def make_paths(tmp_path: Path, registrar: Path | None = None) -> Paths:
    for name in ("fichas", "fichas-ing", "archive"):
        (tmp_path / name).mkdir(exist_ok=True)
    bank = tmp_path / "bank.json"
    questions = [
        {"id": "SYN-T9-001#1", "subtipo": "t9-uno"},
        {"id": "SYN-T9-001#2", "subtipo": None},
        {"id": "SYN-ING-004#1", "subtipo": "morfo-x"},
    ]
    bank.write_text(json.dumps({"schema": 2, "questions": questions, "tests": []}))
    reg = registrar or tmp_path / "unused.py"
    return Paths(
        ledger=tmp_path / "imported.json",
        archive=tmp_path / "archive",
        bank=bank,
        routes={
            "temario": Route(tmp_path / "fichas", reg, tmp_path / "t.xlsx"),
            "ingles": Route(tmp_path / "fichas-ing", reg, tmp_path / "i.xlsx"),
        },
    )


NEVER_OPEN = lambda path: False  # noqa: E731


def ok_runner(calls):
    def run(ficha, registrar, excel):
        calls.append(ficha.name)
        return subprocess.CompletedProcess([], 0, "ok", "")

    return run


def write_results(path: Path, attempts: list[dict]) -> Path:
    path.write_text(json.dumps({"schema": 1, "exported": "x", "attempts": attempts}))
    return path


def cli_args(paths: Paths, results: Path) -> list[str]:
    return [
        "--results", str(results), "--fichas", str(paths.routes["temario"].fichas),
        "--ingles-fichas", str(paths.routes["ingles"].fichas), "--ledger", str(paths.ledger),
        "--archive", str(paths.archive), "--bank", str(paths.bank),
    ]  # fmt: skip


def test_load_results_rejects_bank_and_bad_schema(tmp_path):
    bank_like = tmp_path / "x.json"
    bank_like.write_text(json.dumps({"schema": 1, "questions": [], "tests": []}))
    with pytest.raises(ResultsError, match="not a results file"):
        load_results(bank_like)
    future = tmp_path / "y.json"
    future.write_text(json.dumps({"schema": 2, "attempts": []}))
    with pytest.raises(ResultsError, match="unsupported schema 2"):
        load_results(future)


def test_find_results_picks_newest(tmp_path):
    old = write_results(tmp_path / "results-1.json", [])
    new = write_results(tmp_path / "results-2.json", [])
    (tmp_path / "other.json").write_text("{}")
    os.utime(old, (1, 1))
    assert find_results(tmp_path) == new
    with pytest.raises(FileNotFoundError):
        find_results(tmp_path / "missing")


def test_pending_skips_ledger_and_unfinished_and_sorts():
    a = attempt("a", started="2026-10-14T21:00:00+02:00")
    b = attempt("b", started="2026-10-14T20:00:00+02:00")
    c = attempt("c")
    c["status"] = "in_progress"
    results = {"attempts": [a, b, c]}
    assert [x["uuid"] for x in pending(results, {"zzz"})] == ["b", "a"]
    assert [x["uuid"] for x in pending(results, {"a", "b"})] == []


def test_import_creates_ficha_and_records_ledger(tmp_path):
    paths = make_paths(tmp_path)
    calls: list[str] = []
    registered = import_attempts([attempt()], paths, runner=ok_runner(calls), is_open=NEVER_OPEN)
    assert registered == ["SYN-T9-001"]
    assert calls == ["SYN-T9-001.md"]
    text = (paths.routes["temario"].fichas / "SYN-T9-001.md").read_text(encoding="utf-8")
    assert "respuestas:    C-\n" in text
    assert "subtipos:      1: t9-uno\n" in text
    assert json.loads(paths.ledger.read_text()) == ["u1"]


def test_importing_twice_registers_nothing_the_second_time(tmp_path):
    paths = make_paths(tmp_path)
    results = write_results(tmp_path / "results-1.json", [attempt()])
    calls: list[str] = []
    args = [*cli_args(paths, results), "--yes"]
    assert main(args, runner=ok_runner(calls), is_open=NEVER_OPEN) == 0
    assert main(args, runner=ok_runner(calls), is_open=NEVER_OPEN) == 0
    assert calls == ["SYN-T9-001.md"]
    assert len(list(paths.archive.iterdir())) == 1


def test_prepared_ficha_is_filled_in_place(tmp_path):
    paths = make_paths(tmp_path)
    prepared = paths.routes["temario"].fichas / "SYN-T9-001.md"
    prepared.write_text(
        "---\nid:            SYN-T9-001\nrespuestas:\nclave:         CA\n"
        "subtipos:      1-2: t9-preparado\n---\n\nCuerpo.\n",
        encoding="utf-8",
    )
    import_attempts([attempt()], paths, runner=ok_runner([]), is_open=NEVER_OPEN)
    text = prepared.read_text(encoding="utf-8")
    assert "subtipos:      1-2: t9-preparado\n" in text
    assert text.endswith("\nCuerpo.\n")


def test_second_attempt_gets_suffixed_ficha(tmp_path):
    paths = make_paths(tmp_path)
    calls: list[str] = []
    import_attempts([attempt("u2", number=2)], paths, runner=ok_runner(calls), is_open=NEVER_OPEN)
    assert calls == ["SYN-T9-001-i2.md"]
    assert (
        "id:            SYN-T9-001-i2\n"
        in (paths.routes["temario"].fichas / "SYN-T9-001-i2.md").read_text()
    )


def test_registrar_failure_restores_ficha_and_keeps_earlier_successes(tmp_path):
    paths = make_paths(tmp_path)
    prepared = paths.routes["temario"].fichas / "SYN-T9-002.md"
    original = "---\nid:            SYN-T9-002\nrespuestas:\nclave:         CA\n---\n\nX\n"
    prepared.write_text(original, encoding="utf-8")

    def runner(ficha, registrar, excel):
        code = 1 if ficha.name == "SYN-T9-002.md" else 0
        return subprocess.CompletedProcess([], code, "", "Ya existe SYN-T9-002")

    first = attempt("ok")
    second = attempt("bad", test_id="SYN-T9-002", started="2026-10-14T22:00:00+02:00")
    with pytest.raises(ResultsError, match="SYN-T9-002.*Ya existe"):
        import_attempts([first, second], paths, runner=runner, is_open=NEVER_OPEN)
    assert json.loads(paths.ledger.read_text()) == ["ok"]
    assert prepared.read_text(encoding="utf-8") == original


def test_registrar_failure_removes_new_ficha(tmp_path):
    paths = make_paths(tmp_path)

    def runner(ficha, registrar, excel):
        return subprocess.CompletedProcess([], 1, "", "boom")

    with pytest.raises(ResultsError):
        import_attempts([attempt()], paths, runner=runner, is_open=NEVER_OPEN)
    assert not (paths.routes["temario"].fichas / "SYN-T9-001.md").exists()
    assert not paths.ledger.exists()


def test_cli_declined_confirmation_registers_nothing(tmp_path, monkeypatch):
    paths = make_paths(tmp_path)
    results = write_results(tmp_path / "results-1.json", [attempt()])
    monkeypatch.setattr("builtins.input", lambda prompt: "n")
    calls: list[str] = []
    assert main(cli_args(paths, results), runner=ok_runner(calls), is_open=NEVER_OPEN) == 0
    assert calls == []


@pytest.mark.skipif(
    not (REAL_REGISTRAR.exists() and REAL_EXCEL.exists()), reason="real registrar not here"
)
def test_round_trip_with_real_registrar_on_excel_copy(tmp_path):
    from openpyxl import load_workbook

    excel = tmp_path / "copy.xlsx"
    shutil.copy2(REAL_EXCEL, excel)
    paths = make_paths(tmp_path, registrar=REAL_REGISTRAR)
    paths.routes["temario"] = Route(tmp_path / "fichas", REAL_REGISTRAR, excel)
    test_id = "SYN-T9-901"  # an ID that never exists in the real workbook
    registered = import_attempts([attempt("rt", test_id=test_id)], paths, is_open=NEVER_OPEN)
    assert registered == [test_id]
    ids = [row[0].value for row in load_workbook(excel)["Nivel Test"].iter_rows(min_row=2)]
    assert test_id in ids


def test_attempt_on_an_already_registered_test_gets_the_next_free_suffix(tmp_path):
    paths = make_paths(tmp_path)
    filled = "---\nid:            SYN-T9-001\nrespuestas:    CA\nclave:         CA\n---\n\nX\n"
    (paths.routes["temario"].fichas / "SYN-T9-001.md").write_text(filled, encoding="utf-8")
    (paths.routes["temario"].fichas / "SYN-T9-001-i2.md").write_text(
        filled.replace("SYN-T9-001\n", "SYN-T9-001-i2\n"), encoding="utf-8"
    )
    calls: list[str] = []
    registered = import_attempts(
        [attempt("paper-done")], paths, runner=ok_runner(calls), is_open=NEVER_OPEN
    )
    assert registered == ["SYN-T9-001-i3"]
    text = (paths.routes["temario"].fichas / "SYN-T9-001-i3.md").read_text(encoding="utf-8")
    assert "intento:       3\n" in text
    assert (paths.routes["temario"].fichas / "SYN-T9-001.md").read_text(encoding="utf-8") == filled


def test_app_numbering_restarted_after_reinstall_does_not_collide(tmp_path):
    paths = make_paths(tmp_path)
    import_attempts([attempt("first")], paths, runner=ok_runner([]), is_open=NEVER_OPEN)
    calls: list[str] = []
    import_attempts(
        [attempt("after-reinstall")], paths, runner=ok_runner(calls), is_open=NEVER_OPEN
    )
    assert calls == ["SYN-T9-001-i2.md"]


REAL_ING_REGISTRAR = Path("~/Documents/TheMind/ATC/ingles/99-SCRIPTS/registrar.py").expanduser()
REAL_ING_EXCEL = Path(
    "~/Documents/controlador/Primer Examen/Ingles/Ingles_Maestro_v1.xlsx"
).expanduser()


def test_english_attempt_routes_to_english_fichas_and_registrar(tmp_path):
    paths = make_paths(tmp_path)
    seen = []

    def runner(ficha, registrar, excel):
        seen.append((ficha.parent.name, excel.name))
        return subprocess.CompletedProcess([], 0, "", "")

    a = attempt("en", test_id="SYN-ING-004")
    a["subject"] = "ingles"
    import_attempts([a], paths, runner=runner, is_open=NEVER_OPEN)
    assert seen == [("fichas-ing", "i.xlsx")]
    text = (paths.routes["ingles"].fichas / "SYN-ING-004.md").read_text(encoding="utf-8")
    assert "modo:          Mixto\n" in text and "corte:" not in text


def test_old_export_without_subject_routes_by_prefix(tmp_path):
    paths = make_paths(tmp_path)
    import_attempts(
        [attempt("old", test_id="SYN-ING-004")], paths, runner=ok_runner([]), is_open=NEVER_OPEN
    )
    assert (paths.routes["ingles"].fichas / "SYN-ING-004.md").exists()


def test_paper_registered_english_test_gets_i2(tmp_path):
    paths = make_paths(tmp_path)
    filled = "---\nid:            SYN-ING-004\nrespuestas:    CA\nclave:         CA\n---\n"
    (paths.routes["ingles"].fichas / "SYN-ING-004.md").write_text(filled, encoding="utf-8")
    calls: list[str] = []
    import_attempts(
        [attempt("p", test_id="SYN-ING-004")], paths, runner=ok_runner(calls), is_open=NEVER_OPEN
    )
    assert calls == ["SYN-ING-004-i2.md"]


def test_open_workbook_stops_before_any_write(tmp_path):
    paths = make_paths(tmp_path)
    with pytest.raises(ResultsError, match="i.xlsx.*abierto en Excel"):
        import_attempts(
            [attempt("e", test_id="SYN-ING-004")],
            paths,
            runner=ok_runner([]),
            is_open=lambda p: p.name == "i.xlsx",
        )
    assert not any(paths.routes["ingles"].fichas.iterdir())


def test_unrelated_open_workbook_does_not_block(tmp_path):
    paths = make_paths(tmp_path)
    registered = import_attempts(
        [attempt()], paths, runner=ok_runner([]), is_open=lambda p: p.name == "i.xlsx"
    )
    assert registered == ["SYN-T9-001"]


def test_workbook_open_detects_excel_lock_file(tmp_path):
    from import_results import workbook_open

    book = tmp_path / "Libro.xlsx"
    book.write_text("x")
    assert workbook_open(book) is False
    (tmp_path / "~$Libro.xlsx").write_text("lock")
    assert workbook_open(book) is True


@pytest.mark.skipif(
    not (REAL_ING_REGISTRAR.exists() and REAL_ING_EXCEL.exists()),
    reason="real English registrar not here",
)
def test_english_round_trip_with_real_registrar_on_excel_copy(tmp_path):
    from openpyxl import load_workbook

    excel = tmp_path / "copy-ing.xlsx"
    shutil.copy2(REAL_ING_EXCEL, excel)
    paths = make_paths(tmp_path)
    paths.routes["ingles"] = Route(tmp_path / "fichas-ing", REAL_ING_REGISTRAR, excel)
    a = attempt("rt-en", test_id="SYN-ING-901")
    a["items"][1]["cause"] = "Aplique mal la regla"
    a["items"][1]["ans"] = "b"
    registered = import_attempts([a], paths, is_open=NEVER_OPEN)
    assert registered == ["SYN-ING-901"]
    ids = [row[0].value for row in load_workbook(excel)["Nivel Test"].iter_rows(min_row=2)]
    assert "SYN-ING-901" in ids
